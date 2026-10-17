import json
import threading
import unittest
import urllib.request
from pathlib import Path

from http.server import ThreadingHTTPServer
from io import BytesIO

from PIL import Image
from pptx.enum.shapes import MSO_SHAPE_TYPE

from render import BriefError, build_pptx, validate_brief
from server import handle_rpc, make_handler

BRIEF = json.loads((Path(__file__).resolve().parents[2] / "decks" / "marketplace-brief.json").read_text())


def _png():
    image = Image.new("RGB", (320, 99), (255, 255, 255))
    buffer = BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def fake_brand():
    return {
        "mcp": "https://brand.example/mcp",
        "version": "0.2.0",
        "logoId": "white-logo",
        "colors": {
            "blue": "#006BD9",
            "ink": "#151435",
            "lime": "#D6FF00",
            "cream": "#F8F0E4",
            "white": "#FFFFFF",
            "muted": "#44435D",
        },
        "logo": _png(),
    }


class RenderTests(unittest.TestCase):
    def test_marketplace_brief_builds(self):
        spec = validate_brief(BRIEF)
        self.assertEqual(len(spec["slides"]), 8)
        self.assertIn("build the deck", spec["slides"][0]["headline"])
        import tempfile

        with tempfile.TemporaryDirectory() as directory:
            destination = Path(directory) / "deck.pptx"
            meta = build_pptx(BRIEF, fake_brand(), destination)
            self.assertEqual(meta["slides"], 8)
            self.assertGreater(destination.stat().st_size, 5000)
            from pptx import Presentation

            presentation = Presentation(destination)
            self.assertEqual(len(presentation.slides), 8)
            text = "\n".join(
                shape.text_frame.text
                for slide in presentation.slides
                for shape in slide.shapes
                if shape.has_text_frame
            )
            self.assertIn("The catalogue that can", text)
            self.assertIn("build the deck", text)
            self.assertIn("We built the MCP that builds the deck.", text)
            self.assertIn("This deck was built with an MCP too.", text)
            self.assertIn("wetland-ai-workspace", text)
            pictures = [
                shape
                for slide in presentation.slides
                for shape in slide.shapes
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
            ]
            self.assertGreaterEqual(len(pictures), 2)

    def test_rejects_a_slide_without_a_headline(self):
        broken = json.loads(json.dumps(BRIEF))
        del broken["slides"][1]["headline"]
        with self.assertRaises(BriefError):
            validate_brief(broken)


class McpTests(unittest.TestCase):
    def test_tool_call_writes_inside_the_root(self):
        import tempfile

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            response = handle_rpc(
                {
                    "jsonrpc": "2.0",
                    "id": 2,
                    "method": "tools/call",
                    "params": {"name": "build_deck", "arguments": {"brief": BRIEF, "outputPath": "out/deck.pptx"}},
                },
                root,
                fake_brand,
            )
            self.assertNotIn("error", response)
            self.assertFalse(response["result"].get("isError"))
            self.assertTrue((root / "out" / "deck.pptx").is_file())
            meta = response["result"]["structuredContent"]
            self.assertEqual(meta["tool"], "build_deck")
            self.assertEqual(meta["brandVersion"], "0.2.0")

    def test_refuses_paths_outside_the_root(self):
        import tempfile

        with tempfile.TemporaryDirectory() as directory:
            response = handle_rpc(
                {
                    "jsonrpc": "2.0",
                    "id": 3,
                    "method": "tools/call",
                    "params": {
                        "name": "build_deck",
                        "arguments": {"brief": BRIEF, "outputPath": "../escape.pptx"},
                    },
                },
                Path(directory),
                fake_brand,
            )
            self.assertTrue(response["result"]["isError"])

    def test_http_round_trip(self):
        import tempfile

        with tempfile.TemporaryDirectory() as directory:
            server = ThreadingHTTPServer(("127.0.0.1", 0), make_handler(Path(directory), fake_brand))
            thread = threading.Thread(target=server.serve_forever, daemon=True)
            thread.start()
            port = server.server_address[1]
            try:
                payload = json.dumps(
                    {
                        "jsonrpc": "2.0",
                        "id": 1,
                        "method": "initialize",
                        "params": {
                            "protocolVersion": "2024-11-05",
                            "capabilities": {},
                            "clientInfo": {"name": "test", "version": "0"},
                        },
                    }
                ).encode()
                request = urllib.request.Request(
                    f"http://127.0.0.1:{port}/mcp",
                    data=payload,
                    headers={"Content-Type": "application/json", "Accept": "application/json"},
                )
                with urllib.request.urlopen(request, timeout=5) as response:
                    body = json.loads(response.read().decode())
                self.assertEqual(body["result"]["serverInfo"]["name"], "wetland-deck-mcp")
            finally:
                server.shutdown()


if __name__ == "__main__":
    unittest.main()
